'''
Проверка правильности чтения таблиц из документа
Для того чтобы предотвратить разрывы контекста
'''
import docx
document = docx.Document("./docs/MKS_UM/HE_МКС-УМ.docx")
print(f"Найдено таблиц верхнего уровня: {len(document.tables)}")
for i, table in enumerate(document.tables):
    print(f"--- Таблица {i} ---")
    for row in table.rows:
        print([c.text for c in row.cells])