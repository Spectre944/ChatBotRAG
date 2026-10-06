"""
Загрузка исходных документов (.docx / .pdf / .txt / .md) и извлечение из
них текста и таблиц.

Таблицы извлекаются ОТДЕЛЬНО от обычного текста и намеренно не проходят
через sentence-based чанкинг — каждая таблица целиком становится одним
атомарным чанком, чтобы строка "параметр | значение" никогда не была
разорвана между разными чанками.
"""
import re
from pathlib import Path


def normalize_text(text: str) -> str:
    """
    Чинит типичные артефакты извлечения текста из PDF:
    - слово, разорванное дефисом на конце строки ("слово-\nпродовження")
    - лишние пробелы и переносы строк между буквами/словами
    """
    text = re.sub(r"-\s*\n\s*", "", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def extract_docx_tables(document) -> list[str]:
    """
    Рекурсивно обходит и вложенные таблицы — обычный document.tables их не
    видит, а вложенные таблицы часто появляются в документах, сконвертированных
    из PDF (там, где в оригинале была сложная сетка со значениями).
    """
    def rows_of(table) -> list[str]:
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            rows.append(" | ".join(cells))
            for cell in row.cells:
                for nested in cell.tables:
                    rows.extend(rows_of(nested))
        return rows

    chunks = []
    for table in document.tables:
        rows = rows_of(table)
        if rows:
            chunks.append("\n".join(rows))
    return chunks


def extract_pdf_tables(doc) -> list[str]:
    """Аналог extract_docx_tables для PDF — через PyMuPDF (page.find_tables())."""
    chunks = []
    for page in doc:
        try:
            found = page.find_tables()
        except Exception:
            continue
        for table in found.tables:
            data = table.extract()
            rows = [" | ".join((cell or "").strip() for cell in row) for row in data]
            rows = [r for r in rows if r.strip(" |")]
            if rows:
                chunks.append("\n".join(rows))
    return chunks


def load_documents(folder: str) -> list[dict]:
    """
    Возвращает список документов вида:
        {"text": <обычный текст>, "tables": [<чанк-таблица>, ...], "filename": ...}
    """
    result = []
    for path in Path(folder).glob("**/*"):
        if path.name.startswith("~$"):
            # Временный lock-файл Word, создаётся пока документ открыт — пропускаем
            continue

        if path.suffix.lower() in (".txt", ".md"):
            text = normalize_text(path.read_text(encoding="utf-8"))
            result.append({"text": text, "tables": [], "filename": path.name})

        elif path.suffix.lower() == ".pdf":
            import fitz  # PyMuPDF — извлекает текст надёжнее pypdf на сложных макетах

            doc = fitz.open(str(path))
            text = "\n".join(page.get_text() for page in doc)
            tables = extract_pdf_tables(doc)
            result.append({"text": normalize_text(text), "tables": tables, "filename": path.name})

        elif path.suffix.lower() == ".docx":
            import docx  # python-docx — XML со структурой текста, надёжнее PDF

            document = docx.Document(str(path))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            text = "\n".join(parts)
            tables = extract_docx_tables(document)
            result.append({"text": normalize_text(text), "tables": tables, "filename": path.name})

    return result
