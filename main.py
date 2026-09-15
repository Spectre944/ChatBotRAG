"""
RAG-чат-бот v2 для документации по приборам.

Что добавлено по сравнению с первой версией:
1. Индекс сохраняется на диск — не пересчитывается при перезапуске
2. Документы по новым приборам добавляются в индекс инкрементально,
   с меткой (метаданными), к какому прибору они относятся
3. Эмбеддинги — многоязычная модель (важно для русской документации)
4. Более строгий промпт + возможность посмотреть, что реально нашёл поиск
   (для отладки "бредовых" ответов)

Установка:
    pip install sentence-transformers faiss-cpu pymupdf python-docx requests
"""

import json
import pickle
import re
from pathlib import Path

import faiss
import numpy as np
from sentence_transformers import SentenceTransformer

INDEX_DIR = Path("./index_data")
INDEX_DIR.mkdir(exist_ok=True)
FAISS_PATH = INDEX_DIR / "faiss.index"
META_PATH = INDEX_DIR / "chunks_meta.pkl"

# Многоязычная модель — важно, если документация на русском.
# Английская all-MiniLM хуже ищет релевантные куски в русских текстах.
EMBED_MODEL = "intfloat/multilingual-e5-base"


# ---------- Загрузка и разбиение документов ----------
def normalize_text(text: str) -> str:
    """
    Чинит типичные артефакты извлечения текста из PDF:
    - слово, разорванное дефисом на конце строки ("слово-\nпродовження" -> "словопродовження")
    - лишние пробелы и переносы строк между буквами/словами
    """
    text = re.sub(r"-\s*\n\s*", "", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def extract_docx_tables(document) -> list[str]:
    """
    Каждая таблица целиком превращается в один атомарный чанк (строки
    внутри неё через перенос строки, ячейки — через " | "). Это гарантирует,
    что связка "параметр -> значение" из одной строки таблицы никогда не
    будет разорвана между разными чанками.

    Рекурсивно обходит и вложенные таблицы — обычный document.tables их не
    видит, а вложенные таблицы часто появляются в документах, сконвертированных
    из PDF (там, где в оригинале была сложная сетка со значениями).
    """
    def rows_of(table) -> list[str]:
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            rows.append(" | ".join(cells))
            for cell in row.cells:
                for nested in cell.tables:
                    rows.extend(rows_of(nested))
        return rows

    chunks = []
    for table in document.tables:
        rows = rows_of(table)
        if rows:
            chunks.append("\n".join(rows))
    return chunks


def extract_pdf_tables(doc) -> list[str]:
    """
    Аналог extract_docx_tables для PDF — использует встроенное определение
    таблиц PyMuPDF (page.find_tables()). Требует PyMuPDF >= 1.23.
    """
    chunks = []
    for page in doc:
        try:
            found = page.find_tables()
        except Exception:
            continue
        for table in found.tables:
            data = table.extract()
            rows = [" | ".join((cell or "").strip() for cell in row) for row in data]
            rows = [r for r in rows if r.strip(" |")]
            if rows:
                chunks.append("\n".join(rows))
    return chunks


def load_documents(folder: str) -> list[dict]:
    """
    Возвращает список документов вида
    {"text": <обычный текст>, "tables": [<чанк-таблица>, ...], "filename": ...}
    Таблицы отделены от обычного текста намеренно — см. extract_docx_tables.
    """
    result = []
    for path in Path(folder).glob("**/*"):
        if path.name.startswith("~$"):
            # Временный lock-файл Word, создаётся пока документ открыт —
            # это не настоящий docx, пропускаем
            continue
        if path.suffix.lower() in [".txt", ".md"]:
            text = normalize_text(path.read_text(encoding="utf-8"))
            result.append({"text": text, "tables": [], "filename": path.name})
        elif path.suffix.lower() == ".pdf":
            import fitz  # PyMuPDF — извлекает текст надёжнее pypdf на сложных макетах
            doc = fitz.open(str(path))
            text = "\n".join(page.get_text() for page in doc)
            tables = extract_pdf_tables(doc)
            result.append({"text": normalize_text(text), "tables": tables, "filename": path.name})
        elif path.suffix.lower() == ".docx":
            # DOCX — это XML с реальной структурой текста, парсится намного
            # надёжнее PDF (нет проблемы "угадать порядок символов")
            import docx
            document = docx.Document(str(path))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            text = "\n".join(parts)
            tables = extract_docx_tables(document)
            result.append({"text": normalize_text(text), "tables": tables, "filename": path.name})
    return result


def chunk_text(text: str, target_size: int = 500, overlap_sentences: int = 1) -> list[str]:
    """
    Режем по предложениям, а не по "сырым" символам — иначе в один чанк
    попадает хвост одной мысли и начало другой (это размывает эмбеддинг
    и портит поиск). Собираем предложения в чанк, пока не наберём
    примерно target_size символов, затем начинаем новый с небольшим
    перехлёстом (overlap_sentences), чтобы не терять связку на границе.
    """
    # Разбиение на предложения по украинской/русской пунктуации
    sentences = re.split(r"(?<=[.!?])\s+", text)
    sentences = [s.strip() for s in sentences if s.strip()]

    chunks = []
    current: list[str] = []
    current_len = 0

    for sentence in sentences:
        current.append(sentence)
        current_len += len(sentence)
        if current_len >= target_size:
            chunks.append(" ".join(current))
            current = current[-overlap_sentences:] if overlap_sentences else []
            current_len = sum(len(s) for s in current)

    if current:
        chunks.append(" ".join(current))

    return chunks


# ---------- Индекс с метаданными и персистентностью ----------
class DocumentIndex:
    def __init__(self, model_name: str = EMBED_MODEL):
        self.embedder = SentenceTransformer(model_name)
        self.chunks: list[str] = []
        self.metadata: list[dict] = []  # {"device": ..., "source_file": ...}
        self.index = None

    def exists_on_disk(self) -> bool:
        return FAISS_PATH.exists() and META_PATH.exists()

    def load(self):
        self.index = faiss.read_index(str(FAISS_PATH))
        with open(META_PATH, "rb") as f:
            data = pickle.load(f)
            self.chunks = data["chunks"]
            self.metadata = data["metadata"]
        print(f"Загружен индекс: {len(self.chunks)} чанков")

    def save(self):
        faiss.write_index(self.index, str(FAISS_PATH))
        with open(META_PATH, "wb") as f:
            pickle.dump({"chunks": self.chunks, "metadata": self.metadata}, f)

    def add_documents(self, folder: str, device_name: str):
        """Добавляет документацию по конкретному прибору в индекс (не пересоздавая его)."""
        docs = load_documents(folder)
        new_chunks = []
        new_meta = []
        for doc in docs:
            for chunk in chunk_text(doc["text"]):
                new_chunks.append(chunk)
                new_meta.append({"device": device_name, "source_file": doc["filename"], "type": "text"})

            # Таблицы — отдельными атомарными чанками, без прогона через
            # sentence-based chunk_text, чтобы строка "параметр | значение"
            # никогда не разрывалась между чанками
            for table_chunk in doc["tables"]:
                new_chunks.append(table_chunk)
                new_meta.append({"device": device_name, "source_file": doc["filename"], "type": "table"})

        if not new_chunks:
            print(f"В папке {folder} не нашлось документов")
            return

        # Модель e5 обучена с префиксом "passage: " для индексируемых текстов
        # и "query: " для вопросов — без этого эмбеддинги существенно хуже
        prefixed = [f"passage: {c}" for c in new_chunks]
        embeddings = self.embedder.encode(
            prefixed, show_progress_bar=True, normalize_embeddings=True
        )
        embeddings = np.array(embeddings).astype("float32")

        if self.index is None:
            dim = embeddings.shape[1]
            # IndexFlatIP на нормализованных векторах = косинусное сходство,
            # это метрика, под которую обучена e5 (не L2!)
            self.index = faiss.IndexFlatIP(dim)

        self.index.add(embeddings)
        self.chunks.extend(new_chunks)
        self.metadata.extend(new_meta)
        self.save()
        print(f"Добавлено {len(new_chunks)} чанков по прибору «{device_name}»")

    def search(self, query: str, k: int = 5, device_filter: str | None = None) -> list[dict]:
        q_emb = self.embedder.encode(
            [f"query: {query}"], normalize_embeddings=True
        ).astype("float32")
        search_k = k * 5 if device_filter else k
        _, idx = self.index.search(q_emb, min(search_k, len(self.chunks)))

        results = []
        for i in idx[0]:
            if device_filter and self.metadata[i]["device"] != device_filter:
                continue
            results.append({"text": self.chunks[i], **self.metadata[i]})
            if len(results) >= k:
                break
        return results


# ---------- LLM через Ollama (рекомендуемый вариант) ----------
class OllamaLLM:
    """
    Требует запущенной Ollama (локально: `ollama serve`, модель: `ollama pull mistral`).
    На сервере — тот же код, меняется только base_url.
    """
    def __init__(self, base_url: str = "http://localhost:11434", model_name: str = "mistral"):
        import requests
        self.requests = requests
        self.base_url = base_url
        self.model_name = model_name

    def answer(self, question: str, context_chunks: list[dict]) -> str:
        context = "\n---\n".join(c["text"] for c in context_chunks)
        prompt = build_prompt(question, context)
        return self._generate(prompt)

    def stream_answer(self, question: str, context_chunks: list[dict]):
        """Генератор токенов — общий интерфейс с ClaudeLLM.stream_answer, чтобы app.py не зависел от того, какой backend используется."""
        context = "\n---\n".join(c["text"] for c in context_chunks)
        prompt = build_prompt(question, context)
        resp = self.requests.post(
            f"{self.base_url}/api/generate",
            json={
                "model": self.model_name,
                "prompt": prompt,
                "stream": True,
                "options": {"temperature": 0.2, "repeat_penalty": 1.1},
            },
            stream=True,
        )
        for line in resp.iter_lines():
            if not line:
                continue
            chunk = json.loads(line)
            token = chunk.get("response", "")
            if token:
                yield token
            if chunk.get("done"):
                break

    def raw_generate(self, prompt: str, max_tokens: int = 150) -> str:
        """Обычный вызов модели без строгого RAG-промпта — для вспомогательных задач вроде переформулировки вопроса."""
        return self._generate(prompt, max_tokens=max_tokens, repeat_penalty=1.3)

    def _generate(self, prompt: str, max_tokens: int = 512, repeat_penalty: float = 1.1) -> str:
        resp = self.requests.post(
            f"{self.base_url}/api/generate",
            json={
                "model": self.model_name,
                "prompt": prompt,
                "stream": False,
                "options": {
                    "temperature": 0.2,
                    "num_predict": max_tokens,  # жёсткий потолок длины — не даём модели уйти в зацикливание
                    "repeat_penalty": repeat_penalty,  # штрафует за повторение одинаковых паттернов
                },
            },
        )
        return resp.json()["response"].strip()


GLOSSARY_PATH = Path("./glossary.json")


def load_glossary() -> dict[str, list[str]]:
    if not GLOSSARY_PATH.exists():
        return {}
    with open(GLOSSARY_PATH, "r", encoding="utf-8") as f:
        data = json.load(f)
    data.pop("_readme", None)
    return data


def expand_query_with_glossary(question: str, glossary: dict[str, list[str]]) -> list[str]:
    """
    Быстрая и предсказуемая замена бытовых терминов на официальные из
    словаря — без обращения к LLM. Если в вопросе встречается ключ словаря,
    генерируем вариант вопроса с подставленным официальным термином.
    """
    variants = []
    q_lower = question.lower()
    for term, synonyms in glossary.items():
        if term.lower() in q_lower:
            for syn in synonyms:
                pattern = re.compile(re.escape(term), re.IGNORECASE)
                variants.append(pattern.sub(syn, question))
    return variants


def expand_query(question: str, llm: "OllamaLLM") -> list[str]:
    """
    Пользователь спрашивает бытовыми словами ("сбросить дозу"), а в документации
    может использоваться другая, официальная терминология ("обнулити АЕД").
    Просим LLM предложить альтернативные формулировки и ищем сразу по всем —
    это заметно повышает шанс попасть в нужный раздел документации.
    """
    prompt = f"""Переформулируй вопрос ниже РОВНО 2 альтернативными способами,
используя вероятные официальные технические термины, которые могли бы
использоваться в инструкции к прибору (например, вместо "сбросить дозу"
может быть "обнулити АЕД" или "обнуление накопленной дозы").

Правила:
- Выведи ТОЛЬКО 2 строки, без нумерации и пояснений
- Формулировки должны быть на том же языке, что и вопрос ниже
- Не добавляй ничего, кроме самих 2 формулировок

Вопрос: {question}"""
    raw = llm.raw_generate(prompt, max_tokens=150)
    lines = [line.strip(" -.") for line in raw.split("\n") if line.strip()]
    # На случай, если модель всё равно проигнорирует просьбу и попытается
    # выдать десятки строк — жёстко обрезаем на уровне кода, а не доверяем модели
    lines = [re.sub(r"^\d+\.\s*", "", line) for line in lines][:3]
    return [question] + lines


class ClaudeLLM:
    """
    Использует Claude API вместо локальной/арендованной модели.
    Важно: retrieval (поиск чанков) остаётся полностью локальным —
    Claude API получает уже найденный контекст как обычный текст в
    промпте, у него нет доступа к вашему индексу или файлам.
    claude-haiku-4-5-20251001
    claude-sonnet-5
    """
    def __init__(self, model: str = "claude-haiku-4-5-20251001", api_key: str | None = None):
        import anthropic
        import os
        self.client = anthropic.Anthropic(api_key=api_key or os.environ["ANTHROPIC_API_KEY"])
        self.model = model
        self.system_prompt = (
            "Ты — технический помощник по документации приборов. "
            "Отвечай ТОЛЬКО на основе переданного контекста. Если в контексте "
            "нет ответа — прямо скажи, что информации нет, не выдумывай."
        )

    def answer(self, question: str, context_chunks: list[dict]) -> str:
        context = "\n---\n".join(c["text"] for c in context_chunks)
        resp = self.client.messages.create(
            model=self.model,
            max_tokens=1024,
            system=self.system_prompt,
            messages=[{"role": "user", "content": f"Контекст:\n{context}\n\nВопрос: {question}"}],
        )
        return resp.content[0].text

    def stream_answer(self, question: str, context_chunks: list[dict]):
        """Генератор токенов — тот же интерфейс, что у OllamaLLM.stream_answer."""
        context = "\n---\n".join(c["text"] for c in context_chunks)
        with self.client.messages.stream(
            model=self.model,
            max_tokens=1024,
            system=self.system_prompt,
            messages=[{"role": "user", "content": f"Контекст:\n{context}\n\nВопрос: {question}"}],
        ) as stream:
            for token in stream.text_stream:
                yield token

    def raw_generate(self, prompt: str, max_tokens: int = 150) -> str:
        """Для вспомогательных задач вроде expand_query."""
        resp = self.client.messages.create(
            model=self.model,
            max_tokens=max_tokens,
            messages=[{"role": "user", "content": prompt}],
        )
        return resp.content[0].text


def build_prompt(question: str, context: str) -> str:
    return f"""Ты — технический помощник по документации приборов.
СТРОГО следуй правилам:
- Отвечай ТОЛЬКО на основе контекста ниже.
- Если в контексте нет ответа — прямо скажи "В документации нет информации по этому вопросу", не выдумывай.
- Не додумывай технические характеристики, которых нет в тексте.

Контекст:
{context}

Вопрос: {question}
Ответ:"""


# ---------- Основной сценарий ----------
def main():
    index = DocumentIndex()

    if index.exists_on_disk():
        index.load()
    else:
        print("Индекс не найден, создаю новый и добавляю документы...")
        index.add_documents("./docs/MKS_UM", device_name="МКС-УМ")

    glossary = load_glossary()
    llm = OllamaLLM(model_name="mistral")

    print("\nГотово! Команды: 'exit' — выход, 'add <папка> <прибор>' — добавить документацию")
    while True:
        question = input("\nВопрос: ").strip()
        if question.lower() == "exit":
            break
        if question.lower().startswith("add "):
            _, folder, device = question.split(maxsplit=2)
            index.add_documents(folder, device)
            continue

        # Сначала пробуем быстрый и предсказуемый словарь терминов.
        # К LLM обращаемся только если словарь ничего не подсказал —
        # это и быстрее, и надёжнее (см. проблему с зацикливанием модели).
        glossary_variants = expand_query_with_glossary(question, glossary)
        if glossary_variants:
            queries = [question] + glossary_variants
        else:
            queries = expand_query(question, llm)
        seen_texts = set()
        results = []
        for q in queries:
            for r in index.search(q, k=3):
                key = r["text"][:80]
                if key not in seen_texts:
                    seen_texts.add(key)
                    results.append(r)
        results = results[:5]

        # Отладка: если ответы "бредовые" — сначала смотрим сюда,
        # действительно ли найденные куски релевантны вопросу
        print(f"\n[Формулировки поиска]: {queries}")
        print("[Найденные фрагменты]")
        for r in results:
            print(f"  ({r['device']} / {r['source_file']}): {r['text'][:100]}...")

        answer = llm.answer(question, results)
        print(f"\nОтвет: {answer}")


if __name__ == "__main__":
    main()